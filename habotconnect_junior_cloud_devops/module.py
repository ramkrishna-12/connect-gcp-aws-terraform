from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

base = Path("/mnt/data/habotconnect_junior_cloud_devops")

# ---------- Schema mapping workbook ----------
wb = Workbook()
ws = wb.active
ws.title = "Schema Mapping"

headers = [
    "JavaScript Object Notation Field",
    "Data Type",
    "Required",
    "Validation Rule",
    "Minimum Length",
    "Maximum Length",
    "Yes No Logic",
    "Google BigQuery Column",
    "Notes",
]
rows = [
    ["first_name", "String", "Yes", "Non-empty text", 1, 100, "Not Applicable", "first_name", "Student first name"],
    ["last_name", "String", "Yes", "Non-empty text", 1, 100, "Not Applicable", "last_name", "Student last name"],
    ["email", "String", "Yes", "Valid email address", 3, 254, "Not Applicable", "email", "Contact email"],
    ["date_of_birth", "Date", "Yes", "ISO date; cannot be in the future", 10, 10, "Not Applicable", "date_of_birth", "Calendar date"],
    ["region", "String", "Yes", "Must be APAC, EMEA, or AMER", 4, 4, "Not Applicable", "region", "Data access boundary"],
    ["has_learning_difficulty", "String", "Yes", "Must be Yes or No", 2, 3, "Yes or No", "has_learning_difficulty", "Converted to Boolean"],
    ["receives_learning_support", "String", "Yes", "Must be Yes or No", 2, 3, "Yes or No", "receives_learning_support", "Converted to Boolean"],
    ["needs_learning_support_assistant", "String", "Yes", "Must be Yes or No", 2, 3, "Yes or No", "needs_learning_support_assistant", "Converted to Boolean"],
    ["parental_consent", "String", "Yes", "Must be Yes", 2, 3, "Yes or No", "parental_consent", "Consent gate"],
    ["notes", "String", "No", "Optional text", 0, 500, "Not Applicable", "notes", "Operational notes"],
]
ws.append(["Candidate: Ramkrishna Roy Barman"])
ws.append(["Contact information: add actual contact information before submission."])
ws.append([])
ws.append(headers)
for row in rows:
    ws.append(row)

for row in ws.iter_rows():
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if cell.row == 4:
            cell.font = Font(bold=True)
ws.freeze_panes = "A5"
widths = [32, 18, 12, 42, 16, 18, 20, 32, 34]
for i, width in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = width

notes = wb.create_sheet("Validation Rules")
notes.append(["Candidate: Ramkrishna Roy Barman"])
notes.append(["Contact information: add actual contact information before submission."])
notes.append([])
notes.append(["Rule", "Implementation"])
rules = [
    ["Yes No conversion", "Only Yes and No are accepted, ignoring surrounding whitespace and letter case."],
    ["Consent", "Parental consent must resolve to Yes."],
    ["Unknown fields", "Django REST Framework serializer rejects fields not defined by the contract."],
    ["Region", "Only APAC, EMEA, and AMER are accepted."],
    ["Date of birth", "Must be a valid ISO date and cannot be in the future."],
    ["Notes", "Maximum length is 500 characters."],
]
for row in rules:
    notes.append(row)
for row in notes.iter_rows():
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
notes.column_dimensions["A"].width = 28
notes.column_dimensions["B"].width = 100
for c in notes[4]:
    c.font = Font(bold=True)
for sheet in wb.worksheets:
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

xlsx_path = base / "data" / "schema-mapping.xlsx"
wb.save(xlsx_path)

# ---------- PowerPoint ----------
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

def add_title(slide, title, subtitle=None):
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(12.1), Inches(0.7))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    if subtitle:
        sb = slide.shapes.add_textbox(Inches(0.6), Inches(1.0), Inches(12.1), Inches(0.45))
        sp = sb.text_frame.paragraphs[0]
        sp.text = subtitle
        sp.font.size = Pt(14)

def add_bullets(slide, items, x=0.8, y=1.5, w=11.8, h=5.3, size=20):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.size = Pt(size)
        p.space_after = Pt(10)
        p.level = 0
    return box

def add_node(slide, text, x, y, w=2.4, h=0.9):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.text_frame.text = text
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    shape.text_frame.paragraphs[0].font.size = Pt(15)
    return shape

def add_arrow(slide, x1, y1, x2, y2):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.width = Pt(2)
    return line

# Slide 1
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, "HabotConnect Secure Staging Project", "Junior Cloud & DevOps Engineer | GCP / Django / React")
add_bullets(slide, [
    "Candidate: Ramkrishna Roy Barman",
    "Three deliverables: secure Terraform, fail-closed CI/CD, deterministic Django validation",
    "Objective: turn vague failure conditions into enforceable controls",
], y=1.8, size=22)

# Slide 2
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, "Requirements Deconstructed")
add_bullets(slide, [
    "Task 1: provision D0 Raw Landing Google Cloud Storage and D1 Staged/Enforced BigQuery with strict access controls and row-level security.",
    "Task 2: create a fail-closed build gate for formatting, linting, security scanning, and hardcoded secrets.",
    "Task 3: map student onboarding data into deterministic Yes/No logic and validate it with a Django REST Framework serializer.",
    "Submission: structured code, schema workbook, architecture presentation, and failure demonstration.",
], y=1.4, size=19)

# Slide 3 architecture
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, "Architecture Overview")
nodes = [
    ("Developer Commit", 0.5, 2.5),
    ("Fail-Closed\nGitHub Actions", 3.2, 2.5),
    ("GCS\nD0 Raw Landing", 6.0, 1.6),
    ("BigQuery\nD1 Staged/Enforced", 9.0, 1.6),
    ("Django + DCYN\nValidation", 6.0, 3.7),
]
for text, x, y in nodes:
    add_node(slide, text, x, y)
add_arrow(slide, 2.9, 2.95, 3.2, 2.95)
add_arrow(slide, 5.6, 2.95, 6.0, 2.05)
add_arrow(slide, 8.4, 2.05, 9.0, 2.05)
add_arrow(slide, 7.2, 3.7, 7.2, 2.5)
add_bullets(slide, ["Only a successful gate creates the gated artifact.", "BigQuery table is protected by row-level security."], x=0.7, y=5.4, w=12, h=1.2, size=15)

# Slide 4
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, "Terraform Design")
add_bullets(slide, [
    "Google Cloud Storage bucket: private, uniform bucket-level access, public access prevention, versioning.",
    "Conditional object-creation permission is restricted to the raw/ prefix.",
    "BigQuery dataset: D1 Staged/Enforced with a student_onboarding table.",
    "Dataset and table access are separated by role.",
    "Row-level policy exposes only the configured APAC rows to the analytics group.",
], y=1.4, size=19)

# Slide 5
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, "Storage Security Controls")
add_bullets(slide, [
    "Uniform bucket-level access removes competing object-level access control lists.",
    "Public access prevention blocks accidental public exposure.",
    "Versioning supports recovery from accidental overwrite or deletion.",
    "Least privilege: ingestion receives object creation rather than broad administrative access.",
    "IAM condition limits that permission to the raw/ object prefix.",
], y=1.4, size=19)

# Slide 6
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, "BigQuery and Row-Level Security")
add_bullets(slide, [
    "The student_onboarding table provides the physical data boundary required for row-level policy.",
    "Analytics access is granted through a Google Group rather than individual users.",
    "Example staging predicate: region = 'APAC'.",
    "The production business predicate is intentionally isolated because the hiring brief does not define it.",
    "This prevents unrestricted analytics access to all student rows.",
], y=1.4, size=19)

# Slide 7
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, "Fail-Closed CI/CD Gate")
add_bullets(slide, [
    "Terraform fmt -check",
    "Terraform initialization and validation",
    "TFLint",
    "Checkov infrastructure security scan",
    "Ruff lint and format checks",
    "Django/DCYN unit tests",
    "Gitleaks hardcoded-secret scan",
], x=1.0, y=1.35, w=5.5, h=5.5, size=18)
add_bullets(slide, [
    "All checks are prerequisites for the artifact job.",
    "Any failure stops the gate.",
    "No successful gate means no gated artifact.",
    "This is a control, not merely a warning.",
], x=7.0, y=1.35, w=5.2, h=4.5, size=19)

# Slide 8
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, "Hardcoded Secret Protection")
add_bullets(slide, [
    "Gitleaks scans repository history and working content for credential patterns.",
    "The workflow uses a full-history checkout so historical exposure is not silently ignored.",
    "A detected secret fails the security gate.",
    "The build artifact depends on the successful gate.",
    "No real credentials are included in the demonstration.",
], y=1.4, size=19)

# Slide 9
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, "DCYN Schema Mapping")
add_bullets(slide, [
    "Yes/No fields are accepted only as deterministic binary values.",
    "Yes becomes True; No becomes False.",
    "Values such as Maybe are rejected.",
    "Parental consent must resolve to Yes.",
    "The mapping workbook documents field type, requirement status, limits, logic, and BigQuery destination.",
], y=1.4, size=19)

# Slide 10
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, "Django REST Framework Validation")
add_bullets(slide, [
    "Names: 1-100 characters.",
    "Email: valid format, maximum 254 characters.",
    "Date of birth: ISO date and cannot be in the future.",
    "Region: APAC, EMEA, or AMER.",
    "Notes: optional, maximum 500 characters.",
    "Unknown fields are rejected by the serializer contract.",
], y=1.4, size=19)

# Slide 11
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, "Failure Demonstration")
add_bullets(slide, [
    "Scenario A: introduce a Terraform formatting error → Terraform formatting gate fails → artifact job is skipped.",
    "Scenario B: introduce a fake test credential → Gitleaks reports a secret → security gate fails → artifact job is skipped.",
    "Scenario C: submit Maybe for a Yes/No field → serializer rejects the payload.",
    "Scenario D: submit parental consent as No → serializer rejects the payload.",
], y=1.4, size=19)

# Slide 12
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_title(slide, "Final Submission Checklist")
add_bullets(slide, [
    "Terraform code completed and reviewed.",
    "Fail-closed GitHub Actions workflow completed.",
    "Django REST Framework serializer and DCYN library completed.",
    "Schema workbook has Wrap Text enabled.",
    "Presentation is within the 15-slide limit.",
    "Add actual contact information before submission.",
    "Run local checks and demonstrate one failing security gate.",
    "Submit through the hiring project's Google Form by 13 September 2026.",
], y=1.35, size=18)

pptx_path = base / "presentation" / "habotconnect_project.pptx"
prs.save(pptx_path)

# ---------- Zip entire project ----------
zip_path = Path("/mnt/data/habotconnect_junior_cloud_devops_submission.zip")
with ZipFile(zip_path, "w", ZIP_DEFLATED) as z:
    for path in base.rglob("*"):
        if path.is_file():
            z.write(path, path.relative_to(base.parent))

print(f"Created project folder: {base}")
print(f"Created workbook: {xlsx_path}")
print(f"Created presentation: {pptx_path}")
print(f"Created submission ZIP: {zip_path}")
